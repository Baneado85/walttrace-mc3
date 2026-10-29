"""File-grounded indexing and retrieval; no sample answers or question lookup."""
from __future__ import annotations

from collections import Counter
import csv
from dataclasses import asdict, dataclass
import io
import json
import math
from pathlib import Path
import re
import unicodedata
import zipfile
import xml.etree.ElementTree as ET


def normalized(text: str) -> str:
    return re.sub(r'\s+', ' ', unicodedata.normalize('NFKC', text)).strip().casefold()


def tokens(text: str) -> list[str]:
    whole = re.findall(r'[\w]+(?:[-.#][\w]+)*', text.casefold())
    # Search prose can refer to words embedded in snake_case or hyphenated IDs.
    return whole + [part for token in whole for part in re.split(r'[_-]', token)
                    if part and part != token]


@dataclass
class Chunk:
    path: str
    location: str
    text: str
    withdrawn: bool = False


def segment(path: str, location: str, text: str) -> list[Chunk]:
    lines = text.splitlines()
    withdrawn = bool(re.search(r'withdrawn|superseded|obsolete', path, re.I))
    chunks, current = [], ''
    for line in lines:
        if len(current) + len(line) > 1600 and current:
            chunks.append(Chunk(path, location, current, withdrawn))
            current = current[-180:] + '\n'
        current += line + '\n'
    if current.strip():
        chunks.append(Chunk(path, location, current, withdrawn))
    return chunks


def parse_file(file: Path, relative: str, image_reader=None) -> list[Chunk]:
    ext = file.suffix.casefold()
    if file.stat().st_mode & 0o444 == 0:
        raise PermissionError('file has no read permission bits')
    # Actually open the file; never execute source files or decrypt protected PDFs.
    with file.open('rb') as stream:
        data = stream.read()
    if ext in {'.txt', '.log', '.py', '.md'}:
        return segment(relative, 'text', data.decode('utf-8', errors='replace'))
    if ext == '.csv':
        rows = list(csv.reader(io.StringIO(data.decode('utf-8-sig'))))
        if not rows:
            return []
        header = rows[0]
        return [Chunk(relative, f'row {n}', '; '.join(
            f'{header[i] if i < len(header) else i}: {value}'
            for i, value in enumerate(row)) + '\nCSV row: ' + ','.join(row))
            for n, row in enumerate(rows[1:], 2)]
    if ext == '.docx':
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            document = ET.fromstring(archive.read('word/document.xml'))
        paragraphs = [''.join(t.text or '' for t in p.findall('.//w:t', ns))
                      for p in document.findall('.//w:p', ns)]
        return segment(relative, 'paragraphs/tables', '\n'.join(paragraphs))
    if ext == '.xlsx':
        import openpyxl
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        result = []
        try:
            for sheet in workbook:
                rows = list(sheet.iter_rows(values_only=True))
                for start in range(0, len(rows), 12):
                    group = ([rows[0]] if start else []) + rows[start:start + 12]
                    text = '\n'.join(' | '.join('' if c is None else str(c) for c in row)
                                     for row in group)
                    result.extend(segment(relative, f'{sheet.title} rows {start + 1}+', text))
        finally:
            workbook.close()
        return result
    if ext == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError('encrypted PDF skipped; no password attempted')
        result = []
        for n, page in enumerate(reader.pages):
            text = page.extract_text() or ''
            if len(text.strip()) < 40 and image_reader:
                import pypdfium2 as pdfium
                pdf = pdfium.PdfDocument(data)
                try:
                    bitmap = pdf[n].render(scale=1.5)
                    text = image_reader(bitmap.to_pil().convert('RGB'))
                finally:
                    pdf.close()
            result.extend(segment(relative, f'page {n + 1}', text))
        return result
    if ext in {'.png', '.jpg', '.jpeg', '.tif', '.tiff'}:
        if image_reader is None:
            raise RuntimeError('vision backend required for images')
        from PIL import Image
        with Image.open(io.BytesIO(data)) as image:
            return segment(relative, 'image OCR/relations', image_reader(image.convert('RGB')))
    raise ValueError(f'unsupported extension: {ext}')


def build_index(corpus: Path, image_reader=None) -> tuple[list[Chunk], list[dict]]:
    chunks, skipped = [], []
    root = corpus.resolve()
    for file in sorted(root.rglob('*')):
        if not file.is_file():
            continue
        relative = file.relative_to(root).as_posix()
        try:
            if file.is_symlink() or not file.resolve().is_relative_to(root):
                raise ValueError('symlink outside accepted corpus files')
            chunks.extend(parse_file(file, relative, image_reader))
        except Exception as error:
            skipped.append({'path': relative, 'reason': f'{type(error).__name__}: {error}'})
    return chunks, skipped


class Retriever:
    def __init__(self, chunks: list[Chunk]):
        self.chunks = chunks
        self.bags = [Counter(tokens(c.path + ' ' + c.text)) for c in chunks]
        self.df = Counter(t for bag in self.bags for t in bag)
        self.avg = sum(sum(b.values()) for b in self.bags) / max(1, len(chunks))

    def rank(self, query: str, limit: int = 10) -> list[Chunk]:
        terms = set(tokens(query))
        scored = []
        for i, bag in enumerate(self.bags):
            size = sum(bag.values())
            score = 0.0
            for term in terms:
                freq = bag[term]
                if freq:
                    idf = math.log(1 + (len(self.bags) - self.df[term] + .5) / (self.df[term] + .5))
                    score += idf * freq * 2.2 / (freq + 1.2 * (.25 + .75 * size / max(1, self.avg)))
            if self.chunks[i].withdrawn:
                score *= .15
            if score > 0:
                scored.append((score, i))
        scored.sort(reverse=True)
        return [self.chunks[i] for _, i in scored[:limit]]

    def retrieve(self, query: str) -> list[Chunk]:
        initial = self.rank(query, 7)
        # Follow discovered IDs to retrieve bridging documents in multi-hop questions.
        identifiers = set(re.findall(r'\b[A-Z]{2,12}-\d{2,8}\b', '\n'.join(c.text for c in initial[:4])))
        extra = self.rank(query + ' ' + ' '.join(sorted(identifiers)), 10)
        result, seen = [], set()
        for chunk in initial + extra:
            key = (chunk.path, chunk.location, chunk.text)
            if key not in seen:
                result.append(chunk)
                seen.add(key)
        return result[:12]


def validate_answer(payload: dict, evidence: list[Chunk]) -> dict:
    empty = {'answer': '', 'citations': [], 'confidence': 0.0}
    answer = payload.get('answer')
    supports = payload.get('supports')
    if not isinstance(answer, str) or not answer.strip() or not isinstance(supports, list) or not supports:
        return empty
    selected = []
    for support in supports:
        if not isinstance(support, dict):
            return empty
        path = support.get('path')
        # Some generations copy a source object rather than rename its text key.
        # Either spelling still has to match the actual indexed text below.
        quote = support.get('quote', support.get('text'))
        if not isinstance(path, str) or not isinstance(quote, str) or len(quote.strip()) < 2:
            return empty
        matches = [c for c in evidence if c.path == path and normalized(quote) in normalized(c.text)]
        if not matches or all(c.withdrawn for c in matches):
            return empty
        selected.append((path, quote))
    # The value must appear in an actual supporting quotation, never only in the question.
    compact = lambda t: re.sub(r'[\s.·_\-°]', '', normalized(t))
    if not any(compact(answer) in compact(quote) for _, quote in selected):
        return empty
    return {'answer': answer.strip(), 'citations': sorted(set(p for p, _ in selected)),
            'confidence': .8}


def save_index(directory: Path, corpus: Path, chunks: list[Chunk], skipped: list[dict]) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    data = {'corpus': str(corpus.resolve()), 'chunks': [asdict(c) for c in chunks], 'skipped': skipped}
    temp = directory / 'index.json.tmp'
    temp.write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')
    temp.replace(directory / 'index.json')
