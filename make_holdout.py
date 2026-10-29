"""Create NEW synthetic documents, distinct from the organizer's public sample."""
import json
from pathlib import Path
import zipfile
from reportlab.pdfgen import canvas
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont


def main():
    base = Path.cwd() / 'work/mc3-holdout'
    root = base / 'corpus'
    for folder in ['manuals', 'planning', 'parts', 'events', 'engineering', 'visuals', 'empty']:
        (root / folder).mkdir(parents=True, exist_ok=True)
    for name, voltage, note in [('VX80_current.pdf', 72, 'Current approved specification'),
                                ('VX80_WITHDRAWN.pdf', 120, 'WITHDRAWN: do not use')]:
        page = canvas.Canvas(str(root / 'manuals' / name))
        for y, text in [(760, note), (730, 'VX-80 control module'), (700, f'Maximum supply voltage: {voltage} V')]:
            page.drawString(50, y, text)
        page.save()
    doc = Document()
    doc.add_heading('Pilot schedule', 0)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Device'
    table.cell(0, 1).text = 'Customer sampling'
    table.cell(1, 0).text = 'VX-90'
    table.cell(1, 1).text = 'Q2 FY31'
    doc.save(root / 'planning' / 'pilot_schedule.docx')
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Replacement boards'
    sheet.append(['Device', 'Component', 'Part number'])
    sheet.append(['VX-80', 'field-replaceable controller board', 'PN-X3-70'])
    sheet.append(['VX-20', 'legacy controller board', 'PN-K8-11'])
    workbook.save(root / 'parts' / 'service_catalog.xlsx')
    (root / 'events' / 'incident.log').write_text(
        '2028-02-17 VX-80 pressure trip: error E9902; root cause ticket EVT-5521.\n', encoding='utf-8')
    (root / 'events' / 'fixes.csv').write_text(
        'ticket,summary,fixed_in\nEVT-5521,Input sensor excursion under pulsed load,9.8.7\n'
        'EVT-5522,Display refresh fault,9.8.6\n', encoding='utf-8')
    (root / 'engineering' / 'queue.py').write_text(
        '# Increased from 20 after a backlog.\nDEFAULT_QUEUE_TIMEOUT_S = 275\n', encoding='utf-8')
    font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 40)
    for name, lines in [('wiring.png', ['VX-80 WIRING', 'ALERT_N  ->  PIN D22', 'RESET_N  ->  PIN A07']),
                        ('board.jpg', ['ASSET VX-80', 'BOARD REVISION: REV-H7', 'SERIAL: NX-9302'])]:
        image = Image.new('RGB', (1000, 340), 'white')
        draw = ImageDraw.Draw(image)
        for n, text in enumerate(lines):
            draw.text((40, 30 + n * 90), text, font=font, fill='black')
        image.save(root / 'visuals' / name)
    questions = [
        ('What is the maximum supply voltage of the VX-80?', '72', ['manuals/VX80_current.pdf']),
        ('In which quarter does VX-90 enter customer sampling?', 'Q2 FY31', ['planning/pilot_schedule.docx']),
        ('What is the part number of the field-replaceable controller board for VX-80?', 'PN-X3-70', ['parts/service_catalog.xlsx']),
        ('Which firmware version fixed ticket EVT-5521?', '9.8.7', ['events/fixes.csv']),
        ('What error code is logged for the VX-80 pressure trip?', 'E9902', ['events/incident.log']),
        ('What is the default queue timeout in seconds?', '275', ['engineering/queue.py']),
        ('Which pin carries ALERT_N on the VX-80?', 'D22', ['visuals/wiring.png']),
        ('What board revision is printed on the asset label?', 'REV-H7', ['visuals/board.jpg']),
        ('Which firmware release fixed the root cause of the pressure trip in the incident log?', '9.8.7', ['events/incident.log', 'events/fixes.csv']),
        ('What is the unit price of the VX-80 for 500 units?', '', []),
    ]
    rows = [{'n': n, 'query': q, 'expected_answer': a, 'expected_citations': c,
             'answer_aliases': [a + ' V'] if n == 1 else []}
            for n, (q, a, c) in enumerate(questions, 1)]
    (base / 'questions.json').write_text(json.dumps({'queries': rows}, indent=2), encoding='utf-8')
    target = Path('outputs/mc3-holdout.zip')
    with zipfile.ZipFile(target, 'w', zipfile.ZIP_DEFLATED) as archive:
        for file in base.rglob('*'):
            if file.is_file():
                archive.write(file, file.relative_to(base).as_posix())
    print(target)


if __name__ == '__main__':
    main()
